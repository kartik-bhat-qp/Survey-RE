#!/usr/bin/env python3
"""Convert chat-log JSON stored in Excel cells into a simple transcript workbook.

Input: an .xlsx (or .csv) where one column holds a conversation as JSON, e.g.
    [{"role": "assistant", "content": "Hi! How often do you ..."},
     {"role": "user", "content": "Every week."}]

Output: an .xlsx with one row per turn, ready to import into other systems:
    Transcript ID | <other input columns> | Turn | Speaker | Text

Usage:
    python3 chat_json_to_transcript.py interviews.xlsx
    python3 chat_json_to_transcript.py interviews.xlsx -o transcript.xlsx
    python3 chat_json_to_transcript.py interviews.xlsx --column "Conversation" --id-column "Response ID"
    python3 chat_json_to_transcript.py interviews.xlsx --layout text
    python3 chat_json_to_transcript.py interviews.xlsx --speaker-map "assistant=AI Moderator,user=Participant"

Layouts:
    turns      (default) one row per message
    text       one row per conversation, the whole dialogue in a single "Transcript" cell
    questions  one column per interviewer question, one row per conversation.
               Translates to English first, then groups reworded questions together.

    python3 chat_json_to_transcript.py interviews.xlsx --layout questions -o by-question.xlsx
    python3 chat_json_to_transcript.py transcript-english.xlsx --merge-columns -o grouped.xlsx
"""

from __future__ import annotations

import argparse
import ast
import csv
import json
import os
import re
import sys
import urllib.request
from typing import Any, Dict, List, Optional, Sequence, Tuple

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - dependency guard
    sys.stderr.write(
        "openpyxl is required. Install it with:\n"
        "  python3 -m venv .venv && .venv/bin/pip install -r requirements.txt\n"
    )
    raise SystemExit(2)

# Excel stores at most 32,767 characters per cell.
MAX_CELL_CHARS = 32_000
# Cosine similarity at which two English questions share one column.
SAME_QUESTION_THRESHOLD = 0.75
OTHER_FOLLOW_UPS = "Other follow-ups"

COMPARISON_STOP = set(
    "a an the of to for you your that this what how when where who why would do does did "
    "if or and in on at with rather than just most be is are was were it they them their "
    "me my we our about from by as not only also very really truly more bit could can "
    "please tell say have has had".split()
)

DEFAULT_SPEAKERS = {
    "assistant": "Interviewer",
    "ai": "Interviewer",
    "bot": "Interviewer",
    "model": "Interviewer",
    "user": "Respondent",
    "human": "Respondent",
    "system": "System",
    "tool": "Tool",
    "function": "Tool",
}

# Used when a cell was cut off at the Excel limit and is no longer valid JSON.
SALVAGE_RE = re.compile(
    r'"role"\s*:\s*"(?P<role>[^"]+)"\s*,\s*"content"\s*:\s*"(?P<content>(?:[^"\\]|\\.)*)"',
    re.S,
)

Message = Tuple[str, str]


def cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def looks_like_chat(value: Any) -> bool:
    text = cell_text(value)
    return text.startswith("[") and '"role"' in text[:200].replace("'", '"')


def content_to_text(content: Any) -> str:
    """Flatten OpenAI/Anthropic style content (string, parts list, or dict)."""
    if content is None:
        return ""
    if isinstance(content, str):
        return content.strip()
    if isinstance(content, list):
        parts = [content_to_text(part) for part in content]
        return "\n".join(part for part in parts if part)
    if isinstance(content, dict):
        for key in ("text", "content", "value", "transcript"):
            if key in content:
                return content_to_text(content[key])
        return ""
    return str(content).strip()


def parse_messages(raw: str) -> Tuple[List[Message], bool]:
    """Return (messages, salvaged). salvaged=True means the JSON was broken."""
    data: Any = None
    for loader in (json.loads, ast.literal_eval):
        try:
            data = loader(raw)
            break
        except (ValueError, SyntaxError):
            continue

    if isinstance(data, dict):
        data = data.get("messages") or data.get("conversation") or data.get("transcript")

    if isinstance(data, list):
        messages: List[Message] = []
        for item in data:
            if not isinstance(item, dict):
                continue
            role = cell_text(item.get("role") or item.get("speaker") or item.get("author"))
            text = content_to_text(item.get("content", item.get("text")))
            if text:
                messages.append((role, text))
        return messages, False

    salvaged = []
    for match in SALVAGE_RE.finditer(raw):
        try:
            text = json.loads('"{}"'.format(match.group("content")))
        except ValueError:
            text = match.group("content")
        if text.strip():
            salvaged.append((match.group("role"), text.strip()))
    return salvaged, True


def read_table(path: str, sheet: Optional[str]) -> Tuple[List[str], List[List[Any]]]:
    if path.lower().endswith(".csv"):
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
    else:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        workbook.close()

    rows = [row for row in rows if any(cell_text(cell) for cell in row)]
    if not rows:
        return [], []
    headers = [cell_text(cell) or "Column {}".format(index + 1) for index, cell in enumerate(rows[0])]
    body = [list(row) + [None] * (len(headers) - len(row)) for row in rows[1:]]
    return headers, body


def find_header(headers: Sequence[str], name: str) -> int:
    lowered = [header.lower() for header in headers]
    if name.lower() in lowered:
        return lowered.index(name.lower())
    raise SystemExit('Column "{}" not found. Columns: {}'.format(name, ", ".join(headers)))


def detect_chat_column(headers: Sequence[str], rows: Sequence[List[Any]]) -> int:
    best_index, best_hits = -1, 0
    for index in range(len(headers)):
        hits = sum(1 for row in rows[:200] if looks_like_chat(row[index]))
        if hits > best_hits:
            best_index, best_hits = index, hits
    if best_index < 0:
        raise SystemExit(
            'No column with chat JSON found (expected cells like [{"role": ...}]). '
            "Pass --column to choose one."
        )
    return best_index


def parse_speaker_map(value: Optional[str]) -> Dict[str, str]:
    speakers = dict(DEFAULT_SPEAKERS)
    if not value:
        return speakers
    for pair in value.split(","):
        if "=" not in pair:
            raise SystemExit('Bad --speaker-map entry "{}". Use role=Label.'.format(pair))
        role, label = pair.split("=", 1)
        speakers[role.strip().lower()] = label.strip()
    return speakers


def truncate(text: str) -> str:
    return text if len(text) <= MAX_CELL_CHARS else text[: MAX_CELL_CHARS - 1] + "…"


def write_sheet(
    path: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    text_col: int,
    text_from: Optional[int] = None,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transcript"
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F1F5F9")
    for row in rows:
        sheet.append(list(row))

    wrap = Alignment(vertical="top", wrap_text=True)
    first_text = (text_from if text_from is not None else text_col) + 1
    for row_cells in sheet.iter_rows(min_row=2, min_col=first_text, max_col=len(headers)):
        for cell in row_cells:
            cell.alignment = wrap

    for index, header in enumerate(headers, start=1):
        width = 40 if index >= first_text else min(max(len(header) + 4, 12), 30)
        if text_from is None and index == text_col + 1:
            width = 90
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "B2" if text_from is not None else "A2"

    directory = os.path.dirname(os.path.abspath(path))
    os.makedirs(directory, exist_ok=True)
    workbook.save(path)


def comparison_text(label: str) -> str:
    sentences = re.split(r"(?<=[?.!])\s+", label)
    asking = [sentence for sentence in sentences if "?" in sentence]
    question = asking[-1] if asking else label
    question = re.sub(r"^[^?]{0,80}:\s*", "", question)
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", question.lower())).strip()


def comparison_tokens(label: str) -> set:
    return {
        word
        for word in comparison_text(label).split()
        if len(word) > 3 and word not in COMPARISON_STOP
    }


def questions_ask_the_same_thing(a: str, b: str) -> bool:
    """Rewordings of one question, including a short form of a longer wording."""
    text_a, text_b = comparison_text(a), comparison_text(b)
    shorter, longer = (text_a, text_b) if len(text_a) <= len(text_b) else (text_b, text_a)
    if len(shorter) >= 20 and shorter in longer:
        return True
    tokens_a, tokens_b = comparison_tokens(a), comparison_tokens(b)
    small, big = (tokens_a, tokens_b) if len(tokens_a) <= len(tokens_b) else (tokens_b, tokens_a)
    if len(small) < 2 or not small.issubset(big):
        return False
    return len(small) >= 3 or len(small) / len(big) >= 0.66


def load_api_key() -> str:
    key = os.environ.get("OPENAI_API_KEY", "").strip()
    if key:
        return key
    for path in (".env.local", os.path.join(os.path.dirname(__file__), "..", ".env.local")):
        if not os.path.isfile(path):
            continue
        with open(path, "r", encoding="utf-8") as handle:
            for line in handle:
                if line.startswith("OPENAI_API_KEY="):
                    return line.split("=", 1)[1].strip().strip("'\"")
    raise SystemExit(
        "Translating and grouping questions needs an OpenAI API key in OPENAI_API_KEY or .env.local."
    )


def openai_json(path: str, payload: Dict[str, Any], key: str) -> Dict[str, Any]:
    request = urllib.request.Request(
        "https://api.openai.com/v1/{}".format(path),
        data=json.dumps(payload).encode(),
        headers={"Authorization": "Bearer {}".format(key), "Content-Type": "application/json"},
    )
    with urllib.request.urlopen(request) as response:
        return json.load(response)


def embed_texts(texts: Sequence[str], key: str) -> List[List[float]]:
    vectors: List[Optional[List[float]]] = [None] * len(texts)
    for start in range(0, len(texts), 500):
        chunk = list(texts[start : start + 500])
        data = openai_json(
            "embeddings",
            {"model": "text-embedding-3-small", "input": chunk, "dimensions": 256},
            key,
        )
        for item in data["data"]:
            vector = item["embedding"]
            norm = sum(value * value for value in vector) ** 0.5 or 1
            vectors[start + item["index"]] = [value / norm for value in vector]
    return [vector or [] for vector in vectors]


def cosine(a: Sequence[float], b: Sequence[float]) -> float:
    return sum(left * right for left, right in zip(a, b))


def group_similar_questions(
    labels: Sequence[str], vectors: Sequence[Sequence[float]], weights: Sequence[int]
) -> List[List[int]]:
    """Groups question indexes. The most common wording seeds each group."""
    order = sorted(range(len(labels)), key=lambda index: -weights[index])
    groups: List[Dict[str, Any]] = []
    for index in order:
        best: Optional[Dict[str, Any]] = None
        best_score = SAME_QUESTION_THRESHOLD
        for group in groups:
            if questions_ask_the_same_thing(labels[index], group["label"]):
                best = group
                break
            score = cosine(vectors[index], group["vector"])
            if score >= best_score:
                best, best_score = group, score
        if best is None:
            groups.append({"label": labels[index], "vector": vectors[index], "members": [index]})
        else:
            best["members"].append(index)
    return [group["members"] for group in groups]


def looks_english(text: str) -> bool:
    if re.search(r"[æøåäöüß]", text, re.I):
        return False
    words = re.findall(r"[a-z']+", text.lower())
    english = sum(word in {"the", "and", "you", "your", "to", "what", "would", "that", "this"} for word in words)
    foreign = sum(word in {"und", "ich", "sie", "nicht", "hvad", "ikke", "det"} for word in words)
    return english > 0 and english >= foreign * 2


def translate_to_english(texts: Sequence[str], key: str) -> Dict[str, str]:
    pending = [text for text in dict.fromkeys(texts) if text and not looks_english(text)]
    translated = {text: text for text in texts}
    for start in range(0, len(pending), 40):
        batch = pending[start : start + 40]
        data = openai_json(
            "chat/completions",
            {
                "model": "gpt-4o-mini",
                "temperature": 0,
                "response_format": {"type": "json_object"},
                "messages": [
                    {
                        "role": "system",
                        "content": 'Translate each item into natural English. If it is already English, return it unchanged. Reply with JSON {"t": [...]} with one string per item, in order.',
                    },
                    {"role": "user", "content": json.dumps({"items": batch})},
                ],
            },
            key,
        )
        content = data["choices"][0]["message"]["content"]
        values = json.loads(content).get("t", [])
        for text, value in zip(batch, values):
            if isinstance(value, str):
                translated[text] = value
        print("Translated {}/{}".format(min(start + len(batch), len(pending)), len(pending)), file=sys.stderr)
    return translated


def write_question_workbook(
    records: Sequence[Dict[str, Any]],
    extra_headers: Sequence[str],
    output: str,
    translate: bool,
) -> None:
    if translate:
        texts = [text for record in records for _, text in record["messages"]]
        english = translate_to_english(texts, load_api_key())
        for record in records:
            record["messages"] = [(role, english.get(text, text)) for role, text in record["messages"]]

    interviewer = {"assistant", "ai", "bot", "model", "interviewer"}
    asked_by_record: List[List[Tuple[str, str]]] = []
    label_counts: Dict[str, int] = {}
    for record in records:
        asked: List[Tuple[str, str]] = []
        answers: List[str] = []
        label = ""
        for role, text in record["messages"]:
            if role.lower() in interviewer and "?" in text:
                if label:
                    asked.append((label, "\n".join(answers)))
                sentences = re.split(r"(?<=[?.!])\s+", text)
                asking = [sentence for sentence in sentences if "?" in sentence]
                label = re.sub(r"^[^?]{0,80}:\s*", "", asking[-1] if asking else text).strip()
                answers = []
            elif label and role.lower() not in {"system", "tool", "function"}:
                answers.append(text)
        if label:
            asked.append((label, "\n".join(answers)))
        for question, _answer in asked:
            label_counts[question] = label_counts.get(question, 0) + 1
        asked_by_record.append(asked)

    labels = list(label_counts)
    if not labels:
        raise SystemExit("No interviewer questions found.")
    print("Grouping {} distinct questions…".format(len(labels)), file=sys.stderr)
    vectors = embed_texts(labels, load_api_key())
    groups = group_similar_questions(labels, vectors, [label_counts[label] for label in labels])
    label_column: Dict[str, int] = {}
    column_labels: List[str] = []
    min_count = 1 if len(records) < 2 else max(2, -(-len(records) // 100))
    for members in groups:
        total = sum(label_counts[labels[index]] for index in members)
        if total < min_count:
            continue
        column = len(column_labels)
        column_labels.append(labels[max(members, key=lambda index: label_counts[labels[index]])])
        for index in members:
            label_column[labels[index]] = column

    headers = ["Transcript ID", *extra_headers, *column_labels, OTHER_FOLLOW_UPS]
    rows: List[List[Any]] = []
    for record, asked in zip(records, asked_by_record):
        cells = [""] * len(column_labels)
        follow_ups: List[str] = []
        for question, answer in asked:
            column = label_column.get(question)
            if column is None:
                follow_ups.append("Q: {}\nA: {}".format(question, answer or "—"))
            else:
                cells[column] = "{}\n{}".format(cells[column], answer) if cells[column] else answer
        rows.append([record["id"], *record["extras"], *cells, "\n\n".join(follow_ups)])
    write_sheet(output, headers, rows, text_col=0, text_from=1 + len(extra_headers))
    print(
        "Wrote {} — {} conversation(s), {} question column(s)".format(
            os.path.abspath(output), len(records), len(column_labels)
        )
    )


def merge_question_columns(path: str, output: str) -> int:
    """Collapses reworded question columns of an already-translated workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    table = [list(row) for row in worksheet.iter_rows(values_only=True)]
    workbook.close()
    headers = [cell_text(cell) for cell in table[0]]
    body = table[1:]
    question_indexes = [
        index
        for index, header in enumerate(headers)
        if "?" in header and header != OTHER_FOLLOW_UPS
    ]
    if len(question_indexes) < 2:
        raise SystemExit("No question columns to group (expected headers containing '?').")

    labels = [headers[index] for index in question_indexes]
    weights = [
        sum(1 for row in body if cell_text(row[index] if index < len(row) else None))
        for index in question_indexes
    ]
    print("Comparing {} question columns…".format(len(labels)), file=sys.stderr)
    vectors = embed_texts(labels, load_api_key())
    groups = group_similar_questions(labels, vectors, weights)

    kept_columns: List[int] = []
    new_headers = list(headers)
    redirect: Dict[int, int] = {}
    for members in groups:
        members.sort(key=lambda offset: -weights[offset])
        keep = question_indexes[members[0]]
        kept_columns.append(keep)
        new_headers[keep] = labels[members[0]]
        for offset in members[1:]:
            redirect[question_indexes[offset]] = keep

    drop = set(redirect)
    order = [index for index in range(len(headers)) if index not in drop]
    merged_rows: List[List[Any]] = []
    for row in body:
        padded = list(row) + [None] * (len(headers) - len(row))
        for source, target in redirect.items():
            extra = cell_text(padded[source])
            if not extra:
                continue
            current = cell_text(padded[target])
            padded[target] = "{}\n{}".format(current, extra) if current else extra
        merged_rows.append([padded[index] for index in order])

    out_headers = [new_headers[index] for index in order]
    first_question = min(index for index in order if index in set(kept_columns))
    write_sheet(output, out_headers, merged_rows, text_col=0, text_from=order.index(first_question))
    print(
        "Wrote {} — {} question column(s) grouped into {}".format(
            os.path.abspath(output), len(question_indexes), len(groups)
        )
    )
    for members in groups:
        if len(members) < 2:
            continue
        print("  • {}".format(labels[members[0]]))
        for offset in members[1:]:
            print("      {}".format(labels[offset][:110]))
    return 0


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert chat JSON in Excel cells into a simple transcript workbook."
    )
    parser.add_argument("input", help="Input .xlsx or .csv file")
    parser.add_argument("-o", "--output", help="Output .xlsx (default: <input>-transcript.xlsx)")
    parser.add_argument("--sheet", help="Sheet name to read (default: first sheet)")
    parser.add_argument("--column", help="Column holding the chat JSON (default: auto-detect)")
    parser.add_argument(
        "--id-column", help="Column to use as Transcript ID (default: row number, e.g. T0001)"
    )
    parser.add_argument(
        "--layout",
        choices=("turns", "text", "questions"),
        default="turns",
        help="turns = one row per message; text = one row per conversation; "
        "questions = one column per question, grouped after translation",
    )
    parser.add_argument(
        "--merge-columns",
        action="store_true",
        help="Input is an already-translated workbook. Group similar question columns into one.",
    )
    parser.add_argument(
        "--no-translate",
        action="store_true",
        help="With --layout questions, skip translation and only group the questions",
    )
    parser.add_argument(
        "--speaker-map",
        help='Rename roles, e.g. "assistant=AI Moderator,user=Participant"',
    )
    parser.add_argument(
        "--include-system", action="store_true", help="Keep system/tool messages (dropped by default)"
    )
    parser.add_argument(
        "--no-extra-columns",
        action="store_true",
        help="Do not copy the other input columns into the output",
    )
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    if not os.path.isfile(args.input):
        sys.stderr.write("File not found: {}\n".format(args.input))
        return 1

    if args.merge_columns:
        output = args.output or "{}-grouped.xlsx".format(os.path.splitext(args.input)[0])
        return merge_question_columns(args.input, output)

    headers, rows = read_table(args.input, args.sheet)
    if not rows:
        sys.stderr.write("No data rows found in {}\n".format(args.input))
        return 1

    chat_col = find_header(headers, args.column) if args.column else detect_chat_column(headers, rows)
    id_col = find_header(headers, args.id_column) if args.id_column else None
    speakers = parse_speaker_map(args.speaker_map)
    hidden_roles = set() if args.include_system else {"system", "tool", "function"}

    extra_cols = [] if args.no_extra_columns else [
        index for index in range(len(headers)) if index not in (chat_col, id_col)
    ]
    extra_headers = [headers[index] for index in extra_cols]

    out_rows: List[List[Any]] = []
    question_records: List[Dict[str, Any]] = []
    conversations = turns = 0
    skipped: List[int] = []
    salvaged: List[int] = []

    for row_number, row in enumerate(rows, start=2):
        raw = cell_text(row[chat_col])
        if not raw:
            continue
        messages, was_salvaged = parse_messages(raw)
        messages = [(role, text) for role, text in messages if role.lower() not in hidden_roles]
        if not messages:
            skipped.append(row_number)
            continue
        if was_salvaged:
            salvaged.append(row_number)

        conversations += 1
        transcript_id = cell_text(row[id_col]) if id_col is not None else ""
        transcript_id = transcript_id or "T{:04d}".format(conversations)
        extras = [row[index] for index in extra_cols]

        if args.layout == "questions":
            question_records.append({"id": transcript_id, "extras": extras, "messages": messages})
            turns += len(messages)
            continue

        if args.layout == "turns":
            for turn, (role, text) in enumerate(messages, start=1):
                speaker = speakers.get(role.lower(), role or "Unknown")
                out_rows.append([transcript_id, *extras, turn, speaker, truncate(text)])
                turns += 1
        else:
            dialogue = "\n\n".join(
                "{}: {}".format(speakers.get(role.lower(), role or "Unknown"), text)
                for role, text in messages
            )
            out_rows.append([transcript_id, *extras, len(messages), truncate(dialogue)])
            turns += len(messages)

    if args.layout == "questions":
        if not question_records:
            sys.stderr.write("No conversations could be read from column \"{}\".\n".format(headers[chat_col]))
            return 1
        output = args.output or "{}-transcript.xlsx".format(os.path.splitext(args.input)[0])
        write_question_workbook(question_records, extra_headers, output, translate=not args.no_translate)
        return 0

    if not out_rows:
        sys.stderr.write("No conversations could be read from column \"{}\".\n".format(headers[chat_col]))
        return 1

    if args.layout == "turns":
        out_headers = ["Transcript ID", *extra_headers, "Turn", "Speaker", "Text"]
    else:
        out_headers = ["Transcript ID", *extra_headers, "Turns", "Transcript"]

    output = args.output or "{}-transcript.xlsx".format(os.path.splitext(args.input)[0])
    write_sheet(output, out_headers, out_rows, text_col=len(out_headers) - 1)

    print(
        'Wrote {} — {} conversation(s), {} turn(s) from column "{}"'.format(
            os.path.abspath(output), conversations, turns, headers[chat_col]
        )
    )
    if salvaged:
        print(
            "Warning: rows {} had broken/cut-off JSON (often the Excel 32,767 character limit); "
            "recovered what could be read.".format(", ".join(map(str, salvaged)))
        )
    if skipped:
        print("Skipped rows with no readable messages: {}".format(", ".join(map(str, skipped))))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
